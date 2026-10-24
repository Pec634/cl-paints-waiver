"""Bounded file parsing and reviewed, atomic admin booking imports."""
import csv
import io
import json
import secrets
from datetime import datetime, date, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from flask import request, session, render_template, jsonify, abort
from sqlalchemy import Boolean, Date, DateTime, Integer, Numeric

MAX_ROWS = 200
MAX_COLUMNS = 500
MAX_BYTES = 2 * 1024 * 1024
ADDRESS_PARTS = ('property', 'street', 'city', 'state', 'postcode', 'country', 'country_code')
ADDRESS_FIELDS = {kind: tuple(f'{kind}_address_{part}' for part in ADDRESS_PARTS)
                  for kind in ('client', 'event')}
MAX_EVENT_DAYS = 10
EVENT_FIELDS = ('event_date', 'start_time', 'finish_time', 'event_address',
    *ADDRESS_FIELDS['event'], 'event_type', 'theme', 'pitch_fee_required', 'pitch_fee_amount',
    'expected_attendees', 'publicity_type', 'location_type', 'charge_type',
    'revenue_share_percent', 'celebrates_christmas_easter', 'additional_info', 'estimated_cost')
# Import destinations follow the booking form, rather than every database column.
FORM_FIELDS = (
    'client_type', 'title', 'job_title', 'first_name', 'last_name',
    'date_of_birth', 'is_over_18', 'ethnicity', 'ethnicity_detail', 'religion',
    'phone', 'email', 'client_address', *ADDRESS_FIELDS['client'],
    'event_date', 'start_time', 'finish_time',
    'event_address', *ADDRESS_FIELDS['event'], 'event_type', 'theme', 'pitch_fee_required', 'pitch_fee_amount',
    'expected_attendees', 'publicity_type', 'location_type', 'charge_type',
    'celebrates_christmas_easter', 'additional_info', 'revenue_share_percent', 'estimated_cost', 'payment_preference',
    'promo_code', 'signature',
)


def excel_value(value):
    if value is None:
        return ''
    if isinstance(value, time):
        return value.strftime('%H:%M')
    if isinstance(value, datetime):
        return value.isoformat(sep=' ', timespec='seconds')
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def normalize_time(value):
    """Legacy exports store time-only answers as full Excel datetimes."""
    value = str(value).strip()
    parsed = None
    if ':' in value:
        try:
            parsed = datetime.fromisoformat(value).time()
        except ValueError:
            pass
    if parsed is None:
        for pattern in ('%H:%M', '%H:%M:%S', '%I:%M %p', '%I:%M%p', '%I %p'):
            try:
                parsed = datetime.strptime(value.upper(), pattern).time()
                break
            except ValueError:
                pass
    if parsed is None or parsed.second or parsed.microsecond:
        raise ValueError('use HH:MM, HH:MM:00 or a time such as 2:30 PM (whole minutes).')
    return parsed.strftime('%H:%M')


def linked_excel_events(book):
    """Read legacy event-day worksheets using their explicit parent booking ID."""
    linked = {}
    names = {'Date': 'event_date', 'StartTime': 'start_time', 'FinishTime': 'finish_time',
        'ChargeType': 'charge_type', 'PublicityType': 'publicity_type', 'IndoorsOutdoors': 'location_type',
        'EventType': 'event_type', 'Theme': 'theme', 'IfYouKnowTheNumberOfPeopleAttendingEnterHere': 'expected_attendees',
        'DoTheAttendeesCelebrateChristmasEaster': 'celebrates_christmas_easter',
        'PercentageOfEarningsToShare': 'revenue_share_percent', 'AnyFutherAdditionalInformationThatMayBeHelpfulForUs': 'additional_info',
        'DoCLPaintsNeedToPayAPitchFee': 'pitch_fee_required', 'CostOfPitchFee': 'pitch_fee_amount',
        'TotalPayableForThisEvent': 'estimated_cost'}
    names.update({f'Address_{source}': f'event_address_{dest}' for source, dest in
        zip(('Line1', 'Line2', 'City', 'State', 'PostalCode', 'Country', 'CountryCode'), ADDRESS_PARTS)})
    for sheet in book.worksheets[1:]:
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(values_only=True)
        headers = [excel_value(value) for value in next(iterator, ())]
        if not {'ClientBookingRequest_Id', 'EventDetails_Id', 'Date'}.issubset(headers):
            continue
        for index, values in enumerate(iterator, 1):
            if index > MAX_ROWS * MAX_EVENT_DAYS:
                raise ValueError('The event-day worksheet is too large (maximum 2,000 rows).')
            record = dict(zip(headers, map(excel_value, values)))
            parent = record.get('ClientBookingRequest_Id', '').strip()
            if not parent or not record.get('Date', '').strip():
                continue
            event = {dest: record.get(source, '') for source, dest in names.items()}
            for field in ('start_time', 'finish_time'):
                try:
                    event[field] = normalize_time(event[field])
                except ValueError:
                    pass  # Preserve invalid source values for correction in preview.
            linked.setdefault(parent, []).append(event)
            if len(linked[parent]) > MAX_EVENT_DAYS:
                raise ValueError('A booking can include up to 10 event dates.')
    return linked


def parse_file(filename, data):
    if len(data) > MAX_BYTES:
        raise ValueError('Choose a file under 2 MB.')
    suffix = Path(filename).suffix.lower()
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        from zipfile import ZipFile
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * MAX_BYTES:
                raise ValueError('The expanded workbook is too large.')
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            # Excel's used range includes formatting, often thousands of empty
            # rows/columns. Read real cells rather than trusting that range.
            sheet.reset_dimensions()
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if index > 10000:
                    raise ValueError('The worksheet extends beyond 10,000 rows. Copy the booking table into a new worksheet and retry.')
                values = [excel_value(value) for value in row]
                while values and not values[-1].strip():
                    values.pop()
                if not any(value.strip() for value in values):
                    continue
                rows.append(values)
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError('Import no more than 200 bookings at a time.')
            # Empty styled cells never count as fields. Keep genuinely populated
            # columns, including columns with missing headings, for validation.
            populated_columns = sorted({index for row in rows for index, value in enumerate(row) if value.strip()})
            if len(populated_columns) > MAX_COLUMNS:
                raise ValueError(f'The booking table has more than {MAX_COLUMNS} populated columns. Keep only the columns you want to import.')
            rows = [[row[index] if index < len(row) else '' for index in populated_columns] for row in rows]
            linked = linked_excel_events(book)
            if rows and linked and 'ClientBookingRequest_Id' in rows[0]:
                parent_index = rows[0].index('ClientBookingRequest_Id')
                rows[0].append('events')
                for row in rows[1:]:
                    row.append(json.dumps(linked.get(row[parent_index].strip(), [])))
        finally:
            book.close()
        headers, values = (rows[0], rows[1:]) if rows else ([], [])
    elif suffix in ('.csv', '.tsv'):
        text = data.decode('utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel_tab if suffix == '.tsv' else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        headers, values = (rows[0], rows[1:]) if rows else ([], [])
    elif suffix == '.json':
        rows = json.loads(data.decode('utf-8-sig'))
        if isinstance(rows, dict):
            rows = rows.get('bookings')
        if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
            raise ValueError('JSON must contain an array of bookings or a bookings array.')
        headers = list(dict.fromkeys(key for row in rows for key in row))
        values = [[json.dumps(row[key]) if isinstance(row.get(key), (dict, list)) else str(row.get(key, '') if row.get(key) is not None else '') for key in headers] for row in rows]
    else:
        raise ValueError('Choose CSV, TSV, XLSX or JSON.')
    headers = [str(header).strip() for header in headers]
    values = [row for row in values if any(str(value).strip() for value in row)]
    if not headers or not values:
        raise ValueError('The file has no booking rows.')
    if len(headers) > MAX_COLUMNS or len(values) > MAX_ROWS:
        raise ValueError(f'Use no more than {MAX_COLUMNS} columns and {MAX_ROWS} bookings.')
    if any(not header for header in headers) or len(set(headers)) != len(headers):
        raise ValueError('Column headings must be non-empty and unique.')
    if any(len(row) != len(headers) for row in values):
        raise ValueError('Each row must have the same number of columns as the headings.')
    for index, header in enumerate(headers):
        name = ''.join(character.lower() for character in header if character.isalnum()).rstrip('0123456789')
        if name.endswith(('starttime', 'finishtime', 'endtime')):
            for row in values:
                if row[index].strip():
                    try:
                        row[index] = normalize_time(row[index])
                    except ValueError:
                        pass
    # Legacy form exports split addresses into separate columns. Offer the
    # assembled address as a mapping choice without removing the originals.
    for prefix, destination in (
        ('ClientDetails_Address_', 'client_address'),
        ('EventSectionForASingularDate_Address_', 'event_address'),
        ('CheckAvailabilitySingleDate_AddressOfEvent_', 'availability_event_address'),
    ):
        parts = [headers.index(prefix + part) for part in
                 ('Line1', 'Line2', 'City', 'State', 'PostalCode', 'Country')
                 if prefix + part in headers]
        if parts and destination not in headers:
            headers.append(destination)
            for row in values:
                row.append(', '.join(row[index].strip() for index in parts if row[index].strip()))
    return headers, values


def convert_row(row, columns):
    row = dict(row)
    from booking_birth_dates import client_kind, birth_date
    row['client_type'] = client_kind(row.get('client_type') or 'individual')
    born = birth_date(row.get('date_of_birth'), row['client_type'])
    row['date_of_birth'] = born.isoformat() if born else ''
    raw_events = row.pop('events', None)
    if raw_events in (None, '') and row.get('event_schedule'):
        raw_events = row.pop('event_schedule')
    if isinstance(raw_events, str) and raw_events.strip():
        try:
            raw_events = json.loads(raw_events)
        except ValueError:
            raise ValueError('The event schedule must contain valid JSON.')
    if isinstance(raw_events, dict):
        raw_events = raw_events.get('events')
    if raw_events == []:
        raise ValueError('The event schedule must contain at least one event day.')
    if raw_events:
        if not isinstance(raw_events, list) or not 1 <= len(raw_events) <= MAX_EVENT_DAYS:
            raise ValueError('A booking must contain between 1 and 10 event dates.')
        normalized, first = [], None
        for index, event in enumerate(raw_events, 1):
            if not isinstance(event, dict):
                raise ValueError(f'Event day {index}: invalid event details.')
            base = {key: value for key, value in row.items() if key not in EVENT_FIELDS and key not in ('event_schedule', 'additional_event_dates')}
            base.update({key: value for key, value in event.items() if key in EVENT_FIELDS})
            base['event_date'] = event.get('event_date', event.get('date', ''))
            base['single_date'] = True
            try:
                converted = convert_row(base, columns)
            except ValueError as exc:
                raise ValueError(f'Event day {index}: {exc}') from exc
            if first is None:
                first = converted
            item = {key: converted.get(key) for key in EVENT_FIELDS if key in converted and key != 'event_date'}
            item['date'] = converted['event_date'].isoformat()
            item['duration_hours'] = str((datetime.strptime(converted['finish_time'], '%H:%M') - datetime.strptime(converted['start_time'], '%H:%M')).total_seconds() / 3600)
            holiday = converted['celebrates_christmas_easter']
            item['celebrates_christmas_easter'] = None if holiday is None else ('yes' if holiday else 'no')
            for key in ('estimated_cost', 'pitch_fee_amount'):
                raw = event.get(key, '')
                if raw in ('', None):
                    item[key] = '0.00' if key == 'estimated_cost' else None
                else:
                    try:
                        amount = Decimal(str(raw).replace('£', '').replace(',', ''))
                        if not amount.is_finite() or amount < 0 or amount > Decimal('99999999.99'):
                            raise ValueError()
                        item[key] = str(amount.quantize(Decimal('.01')))
                    except (ValueError, InvalidOperation):
                        raise ValueError(f'Event day {index}: {key} must be a valid non-negative amount.')
            share = event.get('revenue_share_percent', '')
            if share not in ('', None):
                try:
                    item['revenue_share_percent'] = int(share)
                    if item['revenue_share_percent'] not in (5, 10, 15, 20, 25):
                        raise ValueError()
                except (ValueError, TypeError):
                    raise ValueError(f'Event day {index}: choose a revenue share of 5, 10, 15, 20 or 25%.')
            normalized.append(item)
        first['single_date'] = len(normalized) == 1
        first['additional_event_dates'] = ', '.join(item['date'] for item in normalized[1:]) or None
        first['event_schedule'] = json.dumps({'same_details': False, 'events': normalized})
        if any(event.get('estimated_cost') not in ('', None) for event in raw_events):
            first['total_event_cost'] = sum(Decimal(item['estimated_cost']) for item in normalized)
        if first['total_event_cost'] > Decimal('99999999.99'):
            raise ValueError('The total event cost exceeds the booking amount limit.')
        return first
    for kind, fields in ADDRESS_FIELDS.items():
        parts = []
        for field in fields:
            value = row.get(field, '')
            if not isinstance(value, str):
                raise ValueError(f'{field}: use a text value.')
            if len(value) > 200:
                raise ValueError(f'{field}: maximum 200 characters.')
            if value.strip():
                parts.append(value.strip())
        if parts:
            row[f'{kind}_address'] = '\n'.join(parts)
    defaults = dict(request_type='booking', client_type='individual', status='Under Review',
        ethnicity='Not provided', religion='Not provided', client_address='', phone='',
        publicity_type='', charge_type='', location_type='', event_type='', theme='',
        payment_preference='', signature='', is_over_18=False, single_date=True,
        pitch_fee_required=False, liability_acknowledged=False, terms_version='import-unverified',
        terms_accepted_at=datetime.min, travel_charge=Decimal('0'), total_event_cost=Decimal('0'))
    result = {}
    for column in columns:
        name = column.name
        value = row.get(name, '')
        if not isinstance(value, (str, int, float, bool)) and value is not None:
            raise ValueError(f'{name}: use a text value.')
        value = str(value).strip() if value is not None else ''
        if name in ('start_time', 'finish_time') and value:
            try:
                value = normalize_time(value)
            except ValueError as exc:
                raise ValueError(f'{name.replace("_", " ")}: {exc}') from exc
        if not value:
            if name in defaults:
                result[name] = defaults[name]
            elif column.nullable:
                result[name] = None
            elif name == 'submitted_at':
                result[name] = datetime.now()
            else:
                raise ValueError(f'{name.replace("_", " ")}: a value is required.')
            continue
        try:
            if isinstance(column.type, Boolean):
                if value.lower() not in ('true', 'false', 'yes', 'no', '1', '0'):
                    raise ValueError('use Yes or No')
                value = value.lower() in ('true', 'yes', '1')
            elif isinstance(column.type, DateTime):
                value = datetime.fromisoformat(value)
            elif isinstance(column.type, Date):
                try:
                    value = datetime.fromisoformat(value).date()
                except ValueError:
                    value = datetime.strptime(value, '%d/%m/%Y').date()
            elif isinstance(column.type, Integer):
                value = int(value)
                if value < 0:
                    raise ValueError('must not be negative')
            elif isinstance(column.type, Numeric):
                value = Decimal(value.replace('£', '').replace(',', ''))
                if not value.is_finite() or value < 0 or value > Decimal('99999999.99'):
                    raise ValueError('use a valid non-negative amount')
                value = value.quantize(Decimal('.01'))
            elif column.type.length and len(value) > column.type.length:
                raise ValueError(f'maximum {column.type.length} characters')
        except (ValueError, InvalidOperation) as exc:
            raise ValueError(f'{name.replace("_", " ")}: {exc}') from exc
        result[name] = value
    for field in ('start_time', 'finish_time'):
        try:
            result[field] = datetime.strptime(result[field], '%H:%M').strftime('%H:%M')
        except ValueError:
            raise ValueError(f'{field}: use HH:MM (24-hour time).')
    if result['finish_time'] <= result['start_time']:
        raise ValueError('Finish time must be after start time.')
    if result['date_of_birth'] and result['date_of_birth'] > datetime.now().date():
        raise ValueError('Date of birth must not be in the future.')
    if result['liability_acknowledged'] and (not result['signature'] or result['terms_accepted_at'] == datetime.min or result['terms_version'] == 'import-unverified'):
        raise ValueError('Consent requires the original signature, terms acceptance timestamp and terms version. Otherwise leave liability acknowledged as No.')
    if result['status'] not in ('Under Review', 'Accepted', 'Declined', 'Cancelled', 'Completed'):
        raise ValueError('Use a recognised booking status.')
    from email_validator import validate_email, EmailNotValidError
    try:
        result['email'] = validate_email(result['email'], check_deliverability=False).normalized
    except EmailNotValidError as exc:
        raise ValueError(f'Email: {exc}') from exc
    # Complex schedules require the existing scheduling flow; never silently discard them.
    if result.get('event_schedule') or result.get('additional_event_dates') or not result['single_date']:
        raise ValueError('Import each event date as a separate booking; clear schedule / additional dates and set single date to Yes.')
    return result


def register(app, db, Booking, admin_required):
    columns = [column for column in Booking.__table__.columns if column.name != 'id']

    @app.route('/admin/bookings/import', methods=['GET', 'POST'])
    @admin_required
    def admin_import_bookings():
        if request.method == 'GET':
            token = session.setdefault('booking_import_token', secrets.token_urlsafe(32))
            return render_template('admin_booking_import.html', fields=list(FORM_FIELDS), event_fields=list(EVENT_FIELDS), token=token)
        if not secrets.compare_digest(session.get('booking_import_token', ''), request.headers.get('X-Import-Token', '')) or not session.get('booking_import_token'):
            abort(400)
        if request.content_length and request.content_length > 4 * MAX_BYTES:
            return jsonify(error='The import request is too large.'), 400
        if 'file' in request.files:
            upload = request.files['file']
            try:
                headers, rows = parse_file(upload.filename or '', upload.stream.read(MAX_BYTES + 1))
                return jsonify(headers=headers, rows=rows)
            except Exception as exc:
                return jsonify(error=str(exc) if isinstance(exc, ValueError) else 'Unable to read this file. Check its format and try again.'), 400
        payload = request.get_json(silent=True) or {}
        if not isinstance(payload, dict):
            return jsonify(error='Provide an object containing booking rows.'), 400
        rows = payload.get('rows')
        if not isinstance(rows, list) or not rows or len(rows) > MAX_ROWS or not all(isinstance(row, dict) for row in rows):
            return jsonify(error='Provide between 1 and 200 booking rows.'), 400
        records, errors, seen = [], [], set()
        for index, row in enumerate(rows, 1):
            try:
                values = convert_row(row, columns)
                events = json.loads(values['event_schedule'])['events'] if values.get('event_schedule') else [dict(date=values['event_date'].isoformat(), start_time=values['start_time'])]
                existing_keys = set()
                for existing in Booking.query.filter_by(email=values['email']).all():
                    existing_keys.add((existing.event_date.isoformat(), existing.start_time))
                    try:
                        saved = json.loads(existing.event_schedule or '{}')
                        for item in saved.get('events', []) if isinstance(saved, dict) else saved:
                            existing_keys.add((item.get('date'), item.get('start_time')))
                    except (ValueError, TypeError, AttributeError):
                        pass
                for event in events:
                    key = (values['email'], event['date'], event['start_time'])
                    if key in seen or (key[1], key[2]) in existing_keys:
                        raise ValueError(f'A booking with this email, date ({key[1]}) and start time already exists.')
                    seen.add(key)
                records.append(Booking(**values))
            except ValueError as exc:
                errors.append(f'Booking {index}: {exc}')
        if errors:
            return jsonify(error='Nothing was imported. Correct these rows and retry.', errors=errors), 400
        try:
            db.session.add_all(records)
            db.session.commit()
        except Exception:
            db.session.rollback()
            app.logger.exception('Booking import failed')
            return jsonify(error='Nothing was imported. The database could not save these bookings.'), 400
        session.pop('booking_import_token', None)
        return jsonify(imported=len(records))
