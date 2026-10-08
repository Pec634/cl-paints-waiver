"""Admin-only, staged client and booking imports with explicit review."""
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import csv
import hashlib
import json
import re
import secrets
import zipfile
from xml.etree.ElementTree import ParseError

from email_validator import validate_email, EmailNotValidError
from flask import abort, flash, redirect, render_template, request, send_file, session, url_for
from sqlalchemy.exc import IntegrityError
from dateutil.relativedelta import relativedelta
from defusedxml.common import DefusedXmlException

CLIENT_FIELDS = [('email', 'Email', True), ('first_name', 'First name', True),
                 ('last_name', 'Surname', True), ('phone', 'Phone', False), ('address', 'Address', False)]
BOOKING_FIELDS = CLIENT_FIELDS + [(name, label, required) for name, label, required in (
    ('date_of_birth', 'Date of birth', True), ('event_date', 'Event date', True),
    ('start_time', 'Start time', True), ('finish_time', 'Finish time', True),
    ('event_address', 'Event venue/address', True), ('event_type', 'Event name/type', True),
    ('status', 'Status', False), ('publicity_type', 'Public/private', False),
    ('client_type', 'Individual/business', False), ('theme', 'Theme', False),
    ('total_event_cost', 'Event estimate (£)', False), ('travel_charge', 'Travel charge (£)', False),
    ('submitted_at', 'Original submission date', False), ('signature', 'Original signature (if recorded)', False),
    ('source_reference', 'Original booking reference', False))]


def define_models(db):
    class RecordImportBatch(db.Model):
        id = db.Column(db.String(64), primary_key=True)
        owner = db.Column(db.String(64), nullable=False)
        kind = db.Column(db.String(20), nullable=False)
        filename = db.Column(db.String(255), nullable=False)
        payload = db.Column(db.Text, nullable=False)
        mapping = db.Column(db.Text, nullable=False, default='{}')
        state = db.Column(db.String(20), nullable=False, default='Pending')
        created_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
        expires_at = db.Column(db.DateTime, nullable=False)
        imported = db.Column(db.Integer, nullable=False, default=0)
        skipped = db.Column(db.Integer, nullable=False, default=0)
    class RecordImportReceipt(db.Model):
        id = db.Column(db.Integer, primary_key=True)
        key = db.Column(db.String(64), nullable=False, unique=True)
        batch_id = db.Column(db.String(64), db.ForeignKey('record_import_batch.id'), nullable=False)
        client_id = db.Column(db.Integer, db.ForeignKey('client_account.id'))
        booking_id = db.Column(db.Integer, db.ForeignKey('booking.id'))
        source_reference = db.Column(db.String(150), nullable=False, default='')
    return RecordImportBatch, RecordImportReceipt


def read_table(filename, content):
    if not content or len(content) > 5_000_000:
        raise ValueError('Choose a non-empty CSV or XLSX file up to 5 MB.')
    if filename.lower().endswith('.csv'):
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValueError('Save the CSV as UTF-8, or use an XLSX workbook.') from None
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        table = []
        try:
            for row in csv.reader(StringIO(text), dialect):
                if any(str(value).strip() for value in row):
                    table.append(row)
                if len(table) > 501:
                    raise ValueError('Import no more than 500 data rows at a time.')
        except csv.Error:
            raise ValueError('This CSV could not be read. Check its quoting and column format.') from None
    elif filename.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 20_000_000:
                    raise ValueError('This workbook expands beyond the import size limit.')
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                if sheet.max_column and sheet.max_column > 60:
                    raise ValueError('Use no more than 60 columns.')
                table = []
                for physical_row, row in enumerate(sheet.iter_rows(), 1):
                    if physical_row > 501:
                        raise ValueError('Import no more than 500 data rows at a time.')
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError('Replace workbook formulas with their values before importing.')
                    values = [cell.value.isoformat() if isinstance(cell.value, (datetime, date, time)) else str(cell.value) if cell.value is not None else '' for cell in row]
                    if any(value.strip() for value in values):
                        table.append(values)
                    if len(table) > 501:
                        raise ValueError('Import no more than 500 data rows at a time.')
            finally:
                workbook.close()
        except ImportError:
            raise ValueError('Excel support requires installing the updated requirements.') from None
        except (zipfile.BadZipFile, KeyError, OSError, ParseError, DefusedXmlException):
            raise ValueError('Choose a valid, unencrypted XLSX workbook.') from None
    else:
        raise ValueError('Use CSV or Excel XLSX. PDF, scanned files and older XLS workbooks are not supported yet.')
    if len(table) < 2:
        raise ValueError('Include a header row and at least one data row.')
    headers = [str(value).strip() for value in table[0]]
    if len(headers) > 60 or any(not value or len(value) > 150 for value in headers) or len(set(headers)) != len(headers):
        raise ValueError('Use up to 60 non-empty, unique column headings, each up to 150 characters.')
    rows = []
    for values in table[1:]:
        if len(values) > len(headers) or any(len(str(value)) > 4000 for value in values):
            raise ValueError('Check the column count and keep each cell under 4,001 characters.')
        rows.append([str(value).strip() for value in values] + [''] * (len(headers) - len(values)))
    return dict(headers=headers, rows=rows)


def register(app, db, models, Account, Booking, current_client, admin_required):
    Batch, Receipt = models

    def owner():
        return hashlib.sha256(session.setdefault('record_import_owner', secrets.token_hex(32)).encode()).hexdigest()

    def csrf():
        return session.setdefault('client_csrf', secrets.token_urlsafe(32))

    def check():
        if not secrets.compare_digest(csrf(), request.form.get('csrf_token', '')):
            abort(400)

    def parse_date(value):
        for format in ('%Y-%m-%d', '%d/%m/%Y', '%Y-%m-%dT%H:%M:%S'):
            try:
                return datetime.strptime(value, format).date()
            except ValueError:
                pass
        raise ValueError('Use dates as YYYY-MM-DD or DD/MM/YYYY.')

    def money(value):
        amount = Decimal(value.replace('£', '').replace(',', '') or '0')
        if not amount.is_finite() or not 0 <= amount <= Decimal('999999.99') or amount != amount.quantize(Decimal('.01')):
            raise ValueError('Costs must be non-negative amounts with at most two decimal places.')
        return str(amount)

    def validate(values, kind):
        for field, label, required in CLIENT_FIELDS if kind == 'clients' else BOOKING_FIELDS:
            if required and not values.get(field):
                raise ValueError(f'{label} is required.')
        try:
            values['email'] = validate_email(values['email'], check_deliverability=False).normalized.casefold()
        except EmailNotValidError:
            raise ValueError('Enter a valid email address.') from None
        limits = dict(first_name=100, last_name=100, phone=50, address=2000, event_address=2000,
                      event_type=150, theme=200, signature=200, source_reference=150)
        if any(len(values.get(key, '')) > limit for key, limit in limits.items()):
            raise ValueError('A name, address, reference or other field exceeds its allowed length.')
        if kind == 'clients':
            return values
        values['date_of_birth'] = parse_date(values['date_of_birth']).isoformat()
        if parse_date(values['date_of_birth']) > date.today():
            raise ValueError('Date of birth cannot be in the future.')
        values['event_date'] = parse_date(values['event_date']).isoformat()
        for key in ('start_time', 'finish_time'):
            try:
                values[key] = time.fromisoformat(values[key]).strftime('%H:%M')
            except ValueError:
                raise ValueError('Use times as HH:MM, such as 10:30.') from None
        if values['finish_time'] <= values['start_time']:
            raise ValueError('The finish time must be after the start time on the same day.')
        for key, default, allowed in (
            ('status', 'Under Review', ('Under Review', 'Accepted', 'Declined', 'Cancelled')),
            ('publicity_type', 'Private', ('Private', 'Public')),
            ('client_type', 'individual', ('individual', 'business'))):
            value = values.get(key) or default
            normalized = next((item for item in allowed if item.casefold() == value.casefold()), None)
            if normalized is None:
                raise ValueError(f'Choose {key.replace("_", " ")}: ' + ', '.join(allowed) + '.')
            values[key] = normalized
        for key in ('total_event_cost', 'travel_charge'):
            values[key] = money(values.get(key, ''))
        values['submitted_at'] = parse_date(values['submitted_at']).isoformat() if values.get('submitted_at') else date.today().isoformat()
        if parse_date(values['submitted_at']) > date.today():
            raise ValueError('Original submission date cannot be in the future.')
        return values

    def identity(values, kind):
        parts = [kind, values['email']]
        if kind == 'bookings':
            parts += [values.get('source_reference', '').casefold()] if values.get('source_reference') else [
                values[key].casefold() for key in ('event_date', 'start_time', 'finish_time', 'event_address', 'event_type')]
        return hashlib.sha256(json.dumps(parts).encode()).hexdigest()

    def preview(batch):
        payload = json.loads(batch.payload)
        mapping = json.loads(batch.mapping)
        fields = CLIENT_FIELDS if batch.kind == 'clients' else BOOKING_FIELDS
        rows, seen = [], set()
        for index, row in enumerate(payload['rows']):
            values = {key: row[mapping[key]] if key in mapping else '' for key, _, _ in fields}
            error, duplicate, key = '', False, ''
            try:
                values = validate(values, batch.kind)
                key = identity(values, batch.kind)
                duplicate = key in seen or Receipt.query.filter_by(key=key).first() is not None
                if batch.kind == 'clients':
                    duplicate = duplicate or Account.query.filter(db.func.lower(Account.email) == values['email']).first() is not None
                else:
                    existing = Booking.query.filter(db.func.lower(Booking.email) == values['email'],
                        Booking.event_date == parse_date(values['event_date']), Booking.start_time == values['start_time'], Booking.finish_time == values['finish_time']).all()
                    duplicate = duplicate or any(b.event_address.strip().casefold() == values['event_address'].casefold() and b.event_type.strip().casefold() == values['event_type'].casefold() for b in existing)
                seen.add(key)
            except (ValueError, InvalidOperation) as problem:
                error = str(problem) if isinstance(problem, ValueError) else 'Enter a valid numeric cost.'
            rows.append(dict(index=index, number=index + 2, values=values, key=key, error=error, duplicate=duplicate))
        return rows

    @app.route('/admin/imports', methods=['GET', 'POST'])
    @admin_required
    def admin_record_imports():
        expired = Batch.query.filter(Batch.expires_at < datetime.utcnow(), Batch.state == 'Pending').all()
        for batch in expired:
            batch.payload = '{}'; batch.mapping = '{}'; batch.state = 'Expired'
        db.session.commit()
        if request.method == 'POST':
            check()
            try:
                kind = request.form.get('kind')
                if kind not in ('clients', 'bookings'):
                    raise ValueError('Choose clients or bookings.')
                upload = request.files.get('file')
                if not upload:
                    raise ValueError('Choose a file to upload.')
                payload = read_table(upload.filename or '', upload.stream.read(5_000_001))
                batch = Batch(id=secrets.token_hex(32), owner=owner(), kind=kind,
                    filename=(upload.filename or '').replace('\\','/').split('/')[-1][:255], payload=json.dumps(payload),
                    expires_at=datetime.utcnow() + timedelta(hours=24))
                fields = CLIENT_FIELDS if kind == 'clients' else BOOKING_FIELDS
                normalized = {re.sub(r'[^a-z0-9]', '', header.lower()): index for index, header in enumerate(payload['headers'])}
                batch.mapping = json.dumps({key: normalized[re.sub(r'[^a-z0-9]', '', key)] for key, _, _ in fields if re.sub(r'[^a-z0-9]', '', key) in normalized})
                db.session.add(batch); db.session.commit()
                return redirect(url_for('admin_record_import_review', batch_id=batch.id))
            except ValueError as problem:
                flash(str(problem), 'error')
        return render_template('admin_imports.html', client_csrf=csrf())

    @app.route('/admin/imports/<batch_id>', methods=['GET', 'POST'])
    @admin_required
    def admin_record_import_review(batch_id):
        batch = Batch.query.filter_by(id=batch_id, owner=owner()).with_for_update().first_or_404()
        if batch.state == 'Expired' or batch.expires_at < datetime.utcnow() and batch.state != 'Complete':
            abort(410)
        fields = CLIENT_FIELDS if batch.kind == 'clients' else BOOKING_FIELDS
        if request.method == 'POST':
            check()
            if batch.state == 'Complete':
                return redirect(url_for('admin_record_import_review', batch_id=batch.id))
            action = request.form.get('action')
            if action == 'map':
                mapping = {}
                for key, _, _ in fields:
                    value = request.form.get(key, '')
                    if value:
                        try:
                            index = int(value)
                        except ValueError:
                            abort(400)
                        if not 0 <= index < len(json.loads(batch.payload)['headers']):
                            abort(400)
                        mapping[key] = index
                if len(set(mapping.values())) != len(mapping):
                    flash('Match each source column to one destination field.', 'error')
                else:
                    batch.mapping = json.dumps(mapping); db.session.commit()
            elif action == 'save':
                if request.form.get('confirm') != 'yes':
                    abort(400)
                rows = preview(batch)  # Recheck duplicates and validation at save time.
                selected = set(request.form.getlist('row', type=int))
                if not selected or selected - {row['index'] for row in rows}:
                    flash('Select at least one valid row to import.', 'error')
                elif any(row['error'] for row in rows if row['index'] in selected):
                    flash('Correct invalid rows in your source file before importing them.', 'error')
                else:
                    try:
                        for row in rows:
                            if row['index'] not in selected:
                                continue
                            if row['duplicate']:
                                batch.skipped += 1; continue
                            values = row['values']
                            if batch.kind == 'clients':
                                item = Account(email=values['email'], first_name=values['first_name'], last_name=values['last_name'], phone=values['phone'], address=values['address'])
                            else:
                                item = Booking(request_type='booking', client_type=values['client_type'],
                                    first_name=values['first_name'], last_name=values['last_name'], email=values['email'], phone=values['phone'],
                                    client_address=values['address'], date_of_birth=parse_date(values['date_of_birth']), is_over_18=parse_date(values['date_of_birth']) <= date.today()-relativedelta(years=18),
                                    ethnicity='Not recorded', religion='Not recorded', event_date=parse_date(values['event_date']), single_date=True,
                                    start_time=values['start_time'], finish_time=values['finish_time'], event_address=values['event_address'],
                                    event_type=values['event_type'], theme=values['theme'], publicity_type=values['publicity_type'], charge_type='Client pays',
                                    location_type='Not recorded', pitch_fee_required=False, payment_preference='Not recorded',
                                    signature=values['signature'], liability_acknowledged=False, terms_accepted_at=datetime.utcnow(), terms_version='imported-unverified',
                                    status=values['status'], submitted_at=datetime.combine(parse_date(values['submitted_at']), time()),
                                    total_event_cost=Decimal(values['total_event_cost']), travel_charge=Decimal(values['travel_charge']),
                                    internal_notes=f'Imported from {batch.filename}, row {row["number"]}. Original reference: {values["source_reference"] or "Not supplied"}. Consent and payment records have not been verified.')
                            db.session.add(item); db.session.flush()
                            db.session.add(Receipt(key=row['key'], batch_id=batch.id, client_id=item.id if batch.kind == 'clients' else None,
                                booking_id=item.id if batch.kind == 'bookings' else None, source_reference=values.get('source_reference', '')))
                            batch.imported += 1
                        batch.state = 'Complete'; batch.payload = '{}'; batch.mapping = '{}'
                        db.session.commit()
                        return redirect(url_for('admin_record_import_review', batch_id=batch.id))
                    except IntegrityError:
                        db.session.rollback()
                        flash('Another import changed these records. Review the refreshed preview and try again.', 'error')
            else:
                abort(400)
        payload = json.loads(batch.payload)
        response = app.make_response(render_template('admin_import_review.html', batch=batch, fields=fields,
            headers=payload.get('headers', []), mapping=json.loads(batch.mapping), rows=preview(batch) if batch.state == 'Pending' else [], client_csrf=csrf()))
        response.headers['Cache-Control'] = 'private, no-store'
        return response

    @app.get('/admin/imports/template/<kind>')
    @admin_required
    def admin_record_import_template(kind):
        if kind not in ('clients', 'bookings'):
            abort(404)
        output = StringIO(); writer = csv.writer(output)
        writer.writerow([key for key, _, _ in CLIENT_FIELDS if kind == 'clients'] if kind == 'clients' else [key for key, _, _ in BOOKING_FIELDS])
        return send_file(BytesIO(output.getvalue().encode('utf-8-sig')), mimetype='text/csv', as_attachment=True, download_name=f'{kind}-import-template.csv')
