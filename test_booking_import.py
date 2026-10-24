import io
import json
import unittest
from datetime import datetime, time

from booking_import import parse_file, convert_row
import test_business_tools as fixtures

main = fixtures.main


class BookingImportTest(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.BusinessToolsTest()
        self.fixture.setUp()
        self.admin = self.fixture.admin
        self.path = '/admin/bookings/import'
        self.admin.get(self.path)
        with self.admin.session_transaction() as session:
            self.headers = {'X-Import-Token': session['booking_import_token']}
        self.row = dict(first_name='Imported', last_name='Client', date_of_birth='01/02/1990',
            email='imported@example.com', event_date='2027-10-25', start_time='10:00',
            finish_time='12:00', event_address='Test venue')

    def tearDown(self):
        self.fixture.tearDown()

    def test_preview_does_not_write_and_save_is_atomic_and_replay_protected(self):
        count = main.Booking.query.count()
        response = self.admin.post(self.path, headers=self.headers, data={'file':
            (io.BytesIO(json.dumps([self.row]).encode()), 'bookings.json')})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(main.Booking.query.count(), count)
        bad = dict(self.row, email='invalid')
        response = self.admin.post(self.path, headers=self.headers, json={'rows': [self.row, bad]})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(main.Booking.query.count(), count)
        self.row['theme'] = 'Edited preview'
        response = self.admin.post(self.path, headers=self.headers, json={'rows': [self.row]})
        self.assertEqual(response.status_code, 200)
        booking = main.Booking.query.filter_by(email=self.row['email']).one()
        self.assertEqual(booking.theme, 'Edited preview')
        self.assertFalse(booking.liability_acknowledged)
        self.assertEqual(booking.signature, '')
        self.assertEqual(booking.terms_accepted_at, datetime.min)
        self.assertEqual(booking.status, 'Under Review')
        self.assertEqual(self.admin.post(self.path, headers=self.headers, json={'rows': [self.row]}).status_code, 400)

    def test_auth_csrf_duplicates_and_invalid_values(self):
        self.assertEqual(main.app.test_client().get(self.path).status_code, 302)
        self.assertEqual(self.admin.post(self.path, json={'rows': [self.row]}).status_code, 400)
        self.assertEqual(self.admin.post(self.path, headers=self.headers, json={'rows': [self.row, self.row]}).status_code, 400)
        columns = [c for c in main.Booking.__table__.columns if c.name != 'id']
        for changes in ({'travel_charge': 'NaN'}, {'finish_time': '09:00'}, {'event_schedule': '{"events": []}'}, {'first_name': ''}, {'is_over_18': 'unknown'}):
            with self.assertRaises(ValueError):
                convert_row(dict(self.row, **changes), columns)

    def test_supported_files_and_bad_headers(self):
        self.assertEqual(parse_file('test.csv', b'first_name;last_name\nJane;Smith')[1], [['Jane', 'Smith']])
        self.assertEqual(parse_file('test.tsv', b'first_name\tlast_name\nJane\tSmith')[0], ['first_name', 'last_name'])
        from openpyxl import Workbook
        book = Workbook(); book.active.append(['first_name']); book.active.append(['Jane'])
        stream = io.BytesIO(); book.save(stream); book.close()
        self.assertEqual(parse_file('test.xlsx', stream.getvalue())[1], [['Jane']])
        with self.assertRaises(ValueError):
            parse_file('test.csv', b'name,name\nJane,Smith')

    def test_excel_ignores_formatted_empty_rows_and_columns(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        book = Workbook()
        sheet = book.active
        sheet.append([])
        sheet.append(['first_name', 'start_time', 'event_date', 'notes'])
        sheet.append(['Jane', time(10, 30), datetime(2027, 10, 25)])
        sheet.append([])
        sheet.append(['John', time(11), datetime(2027, 10, 26), 'Keep this'])
        sheet.cell(500, 150).fill = PatternFill(fill_type='solid', fgColor='FFFF00')
        stream = io.BytesIO(); book.save(stream); book.close()
        headers, rows = parse_file('bookings.xlsx', stream.getvalue())
        self.assertEqual(headers, ['first_name', 'start_time', 'event_date', 'notes'])
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0], ['Jane', '10:30', '2027-10-25 00:00:00', ''])
        self.assertEqual(rows[1][-1], 'Keep this')

    def test_excel_limits_still_apply_to_real_bookings(self):
        from openpyxl import Workbook
        book = Workbook(); book.active.append(['first_name'])
        for index in range(201):
            book.active.append([f'Client {index}'])
        stream = io.BytesIO(); book.save(stream); book.close()
        with self.assertRaisesRegex(ValueError, '200 bookings'):
            parse_file('bookings.xlsx', stream.getvalue())


if __name__ == '__main__':
    unittest.main()
